"""Exportación de FAQ a JSON, Word (.docx) y PDF."""
import io
import json
from datetime import datetime, timezone

from app.models import Faq, Tenant


def to_json(tenant: Tenant, faqs: list[Faq]) -> bytes:
    payload = {
        "empresa": tenant.name,
        "exportado": datetime.now(timezone.utc).isoformat(),
        "total": len(faqs),
        "faqs": [{"pregunta": f.question, "respuesta": f.answer, "tipo": f.chat_type} for f in faqs],
    }
    return json.dumps(payload, ensure_ascii=False, indent=2).encode()


def to_docx(tenant: Tenant, faqs: list[Faq]) -> bytes:
    import docx
    d = docx.Document()
    d.add_heading(f"Preguntas frecuentes · {tenant.name}", level=0)
    d.add_paragraph(f"Exportado: {datetime.now(timezone.utc):%Y-%m-%d %H:%M} UTC · {len(faqs)} preguntas")
    for f in faqs:
        d.add_heading(f.question, level=2)
        d.add_paragraph(f.answer)
        d.add_paragraph(f"Tipo: {'Soporte técnico' if f.chat_type == 'technical' else 'Soporte comercial'}").italic = True
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def to_pdf(tenant: Tenant, faqs: list[Faq]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm)
    styles = getSampleStyleSheet()
    story = [Paragraph(f"Preguntas frecuentes · {tenant.name}", styles["Title"]),
             Paragraph(f"{len(faqs)} preguntas", styles["Normal"]), Spacer(1, 12)]
    for f in faqs:
        story.append(Paragraph(f.question, styles["Heading3"]))
        story.append(Paragraph(f.answer.replace("\n", "<br/>"), styles["Normal"]))
        story.append(Spacer(1, 8))
    doc.build(story)
    return buf.getvalue()


EXPORters = {"json": (to_json, "application/json", "json"),
             "docx": (to_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
             "pdf": (to_pdf, "application/pdf", "pdf")}
